# -*- coding: utf-8 -*-
"""Wedding RSVP app: per-household guest links + couple admin dashboard.

Stack: Flask + SQLite locally, or Postgres (e.g. Neon) when the
DATABASE_URL environment variable is set — see the DB section below.
Run locally:  python app.py   (serves on http://127.0.0.1:5000)
Deploy: see render.yaml / Procfile / README.md
"""
import csv
import io
import json
import os
import re
import secrets
import sqlite3
import string
import sys
from datetime import datetime
from urllib.parse import urlparse

from flask import Flask, jsonify, request, send_from_directory

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("RSVP_DB", os.path.join(BASE_DIR, "wedding.db"))

app = Flask(__name__, static_folder="static", static_url_path="/static")

# ---------------------------------------------------------------- DB ----

DATABASE_URL = os.environ.get("DATABASE_URL", "")
USE_PG = DATABASE_URL.startswith(("postgres://", "postgresql://"))
SEED_PATH = os.path.join(BASE_DIR, "seed_data.json")


class _Row:
    """Row with sqlite3.Row-style access: row["col"], row[0], dict(row)."""

    __slots__ = ("_keys", "_values", "_index")

    def __init__(self, keys, values):
        self._keys = list(keys)
        self._values = tuple(values)
        self._index = {}
        for i, k in enumerate(self._keys):
            self._index.setdefault(k, i)

    def __getitem__(self, key):
        if isinstance(key, int):
            return self._values[key]
        return self._values[self._index[key]]

    def keys(self):
        return list(self._keys)

    def __iter__(self):
        return iter(self._values)

    def __len__(self):
        return len(self._values)


class _PgCursor:
    """Wraps a psycopg cursor so fetchone/fetchall return _Row objects."""

    def __init__(self, cursor):
        self._cursor = cursor

    def _wrap(self, values):
        if values is None:
            return None
        keys = [d.name for d in self._cursor.description]
        return _Row(keys, values)

    def fetchone(self):
        return self._wrap(self._cursor.fetchone())

    def fetchall(self):
        return [self._wrap(v) for v in self._cursor.fetchall()]


class _PgConnection:
    """Minimal sqlite3-connection lookalike backed by psycopg (v3).

    SQL is written once with sqlite-style '?' placeholders; they are
    translated to psycopg '%s' here so the rest of the app is unchanged.
    """

    def __init__(self, dsn):
        import psycopg

        self._conn = psycopg.connect(dsn)

    def execute(self, sql, params=()):
        cur = self._conn.cursor()
        cur.execute(sql.replace("?", "%s"), tuple(params) if params else None)
        return _PgCursor(cur)

    def commit(self):
        self._conn.commit()

    def close(self):
        self._conn.close()


# Postgres schema. Differences from the SQLite DDL below: BIGSERIAL
# instead of INTEGER PRIMARY KEY AUTOINCREMENT, and created/updated
# timestamp defaults rendered as 'YYYY-MM-DD HH:MM:SS' text (UTC) so the
# stored format matches what SQLite's CURRENT_TIMESTAMP produces.
_PG_TS = "(to_char(now() AT TIME ZONE 'UTC', 'YYYY-MM-DD HH24:MI:SS'))"
_PG_DDL = [
    """
    CREATE TABLE IF NOT EXISTS households (
        id BIGSERIAL PRIMARY KEY,
        token TEXT UNIQUE NOT NULL,
        display_name TEXT NOT NULL,
        group_name TEXT DEFAULT '',
        phone TEXT DEFAULT '',
        invited_total INTEGER DEFAULT 0,
        created_at TEXT DEFAULT %s
    )
    """ % _PG_TS,
    """
    CREATE TABLE IF NOT EXISTS responses (
        household_id BIGINT PRIMARY KEY REFERENCES households(id),
        attending TEXT DEFAULT '',
        party_count INTEGER DEFAULT 0,
        overnight TEXT DEFAULT '',
        updated_at TEXT DEFAULT %s
    )
    """ % _PG_TS,
    """
    CREATE TABLE IF NOT EXISTS settings (
        key TEXT PRIMARY KEY,
        value TEXT
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS send_log (
        id BIGSERIAL PRIMARY KEY,
        household_id BIGINT,
        channel TEXT,
        status TEXT,
        detail TEXT,
        created_at TEXT DEFAULT %s
    )
    """ % _PG_TS,
    """
    CREATE TABLE IF NOT EXISTS send_jobs (
        id BIGSERIAL PRIMARY KEY,
        household_id BIGINT REFERENCES households(id),
        kind TEXT,
        to_phone TEXT,
        message TEXT,
        status TEXT DEFAULT 'pending',
        provider_message_id TEXT,
        error TEXT,
        created_at TEXT DEFAULT %s,
        updated_at TEXT DEFAULT %s
    )
    """ % (_PG_TS, _PG_TS),
]


def db():
    if USE_PG:
        return _PgConnection(DATABASE_URL)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    if USE_PG:
        for _stmt in _PG_DDL:
            conn.execute(_stmt)
    else:
        conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS households (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            token TEXT UNIQUE NOT NULL,
            display_name TEXT NOT NULL,
            group_name TEXT DEFAULT '',
            phone TEXT DEFAULT '',
            invited_total INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS responses (
            household_id INTEGER PRIMARY KEY REFERENCES households(id),
            attending TEXT DEFAULT '',
            party_count INTEGER DEFAULT 0,
            overnight TEXT DEFAULT '',
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT
        );
        CREATE TABLE IF NOT EXISTS send_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            household_id INTEGER,
            channel TEXT,
            status TEXT,
            detail TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS send_jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            household_id INTEGER REFERENCES households(id),
            kind TEXT,
            to_phone TEXT,
            message TEXT,
            status TEXT DEFAULT 'pending',
            provider_message_id TEXT,
            error TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        """
    )
    defaults = {
        "couple_names": "שם החתן ושם הכלה",
        "wedding_date": "",
        "wedding_time": "",
        "venue": "",
        "address": "",
        "info_text": "אנחנו מתחתנים! נשמח לחגוג איתכם את היום המיוחד שלנו. אנא אשרו הגעה בקישור האישי שלכם.",
        "packing_list": (
            "רשימת ציוד ללינה (רשימה זמנית, תוחלף בהמשך):\n"
            "• בגדים להחלפה ליום שאחרי\n"
            "• פיג'מה\n"
            "• מברשת שיניים ומוצרי טיפוח\n"
            "• מגבת\n"
            "• נעליים נוחות\n"
            "• בגד ים (אם יש בריכה)\n"
            "• תרופות אישיות\n"
            "• מטען לטלפון\n"
            "• קרם הגנה וכובע"
        ),
        "invite_template": "היי {name}! 💍 אנחנו מתחתנים! נשמח שתאשרו הגעה בקישור האישי שלכם: {link}",
        "whatsapp_provider": "wa_link",
        "gateway_url": "",
        "gateway_token": "",
        "gateway_template": "",
        "relay_key": "",
        "zapclaw_api_key": "",
        "zapclaw_from": "",
        "zapclaw_invite_template": "wedding_invitation",
        "zapclaw_packing_template": "wedding_packing_list",
    }
    insert_default = (
        "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT DO NOTHING"
        if USE_PG
        else "INSERT OR IGNORE INTO settings (key, value) VALUES (?, ?)"
    )
    for key, value in defaults.items():
        conn.execute(insert_default, (key, value))
    if USE_PG:
        _seed_pg_if_empty(conn)
    conn.commit()
    conn.close()


def _seed_pg_if_empty(conn):
    """One-time import of seed_data.json into a fresh Postgres database.

    Runs only on Postgres, only while the households table is empty, so
    restarts never duplicate the seed. seed_data.json was generated from
    the last known-good local wedding.db (all households including their
    tokens, plus all settings rows).
    """
    if not os.path.exists(SEED_PATH):
        return
    count = conn.execute("SELECT COUNT(*) FROM households").fetchone()[0]
    if count:
        return
    with open(SEED_PATH, encoding="utf-8") as f:
        seed = json.load(f)
    households = seed.get("households", [])
    for h in households:
        conn.execute(
            """INSERT INTO households
               (id, token, display_name, group_name, phone, invited_total, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (h["id"], h["token"], h["display_name"], h.get("group_name") or "",
             h.get("phone") or "", h.get("invited_total") or 0,
             h.get("created_at")),
        )
    for row in seed.get("settings", []):
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) "
            "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
            (row["key"], row["value"]),
        )
    if households:
        # Explicitly seeded ids do not advance the BIGSERIAL sequence;
        # move it past the seeded max so future inserts cannot collide.
        conn.execute(
            "SELECT setval(pg_get_serial_sequence('households', 'id'), "
            "(SELECT MAX(id) FROM households))"
        )


def get_settings():
    conn = db()
    rows = conn.execute("SELECT key, value FROM settings").fetchall()
    conn.close()
    return {r["key"]: r["value"] for r in rows}


def public_settings():
    s = get_settings()
    return {k: s.get(k, "") for k in (
        "couple_names", "wedding_date", "wedding_time", "venue",
        "address", "info_text", "packing_list",
    )}


def new_token():
    alphabet = string.ascii_letters + string.digits
    return "".join(secrets.choice(alphabet) for _ in range(12))


# ------------------------------------------------------------- phones ----

def normalize_phone(raw):
    """Normalize an Israeli phone number to 972XXXXXXXXX digits, or ''."""
    if raw is None:
        return ""
    digits = re.sub(r"\D", "", str(raw))
    if not digits:
        return ""
    if digits.startswith("972"):
        digits = digits[3:]
    elif digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 9 and digits.startswith("5"):
        return "972" + digits
    if len(digits) == 9:
        return "972" + digits
    if len(digits) == 10 and digits.startswith("05"):
        return "972" + digits[1:]
    return digits


def wa_link(phone, message):
    import urllib.parse
    return "https://wa.me/{}?text={}".format(
        phone, urllib.parse.quote(message)
    )


# --------------------------------------------------------- smart import ----

NAME_KEYS = ["שם", "משפחה", "זוג", "אורח", "מוזמן", "name", "guest", "family"]
PHONE_KEYS = ["טלפון", "נייד", "סלולרי", "פלאפון", "phone", "mobile", "tel"]
KIDS_KEYS = ["ילדים", "ילד", "kids", "children"]
TOTAL_KEYS = ["סה\"כ", "סך הכל", "סהכ", "כמות", "סה״כ", "total", "count"]
GROUP_HINTS = ["משפחה", "חברים", "קבוצה", "group"]


def _cell_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_phone_like(text):
    digits = re.sub(r"\D", "", text)
    return len(digits) >= 9 and len(digits) <= 13 and digits[:1] in "05+" or (
        len(digits) >= 9 and re.match(r"^\+?\d[\d\s\-()]{7,}$", text) is not None
    )


def _looks_like_header(row):
    joined = " ".join(row)
    keys = NAME_KEYS + PHONE_KEYS + TOTAL_KEYS + KIDS_KEYS
    return any(k in joined for k in keys)


def _map_header(row):
    """Return dict of role -> column index based on header cells."""
    mapping = {"names": [], "phone": None, "kids": None, "total": None}
    for idx, cell in enumerate(row):
        low = cell
        if any(k in low for k in PHONE_KEYS):
            mapping["phone"] = idx
        elif any(k in low for k in TOTAL_KEYS):
            mapping["total"] = idx
        elif any(k in low for k in KIDS_KEYS):
            mapping["kids"] = idx
        elif any(k in low for k in NAME_KEYS):
            mapping["names"].append(idx)
    return mapping


def parse_rows(rows):
    """Parse raw spreadsheet rows into household dicts + a report.

    Handles: header rows in Hebrew/English, two name columns (couple),
    group title rows (single cell), phone formats, kids/total columns,
    repeated headers, blank rows, and total mismatches.
    """
    households = []
    report = {"added": 0, "updated": 0, "skipped": 0, "warnings": []}
    mapping = None
    group = ""
    for rno, raw in enumerate(rows, start=1):
        row = [_cell_str(c) for c in raw]
        nonempty = [c for c in row if c]
        if not nonempty:
            continue
        # Group title row: a single meaningful cell, no phone, no digits-only total
        if len(nonempty) == 1 and not _is_phone_like(nonempty[0]):
            group = nonempty[0]
            continue
        if _looks_like_header(row):
            m = _map_header(row)
            if m["names"] or m["phone"] is not None:
                mapping = m
            continue
        # --- data row ---
        names, phone, kids, total = [], "", 0, 0
        if mapping:
            for idx in mapping["names"]:
                if idx < len(row) and row[idx]:
                    names.append(row[idx])
            if mapping["phone"] is not None and mapping["phone"] < len(row):
                phone = normalize_phone(row[mapping["phone"]])
            if mapping["kids"] is not None and mapping["kids"] < len(row):
                try:
                    kids = int(float(row[mapping["kids"]]))
                except ValueError:
                    kids = 0
            if mapping["total"] is not None and mapping["total"] < len(row):
                try:
                    total = int(float(row[mapping["total"]]))
                except ValueError:
                    total = 0
        else:
            # No header seen: heuristic — phone-like cell = phone,
            # first 1-2 text cells = names, last number = total.
            for cell in row:
                if _is_phone_like(cell) and not phone:
                    phone = normalize_phone(cell)
                elif cell and not cell.replace(".", "", 1).isdigit():
                    if len(names) < 2:
                        names.append(cell)
            nums = [c for c in row if c.replace(".", "", 1).isdigit()]
            if nums:
                try:
                    total = int(float(nums[-1]))
                except ValueError:
                    total = 0
        if not names:
            report["skipped"] += 1
            report["warnings"].append("שורה {}: אין שם, דולגה".format(rno))
            continue
        display = " ו־".join(names)
        invited = total if total > 0 else len(names) + kids
        adults = len(names)
        if total > 0 and total != adults + kids:
            report["warnings"].append(
                'שורה {} ({}): סה"כ {} לא תואם {} מבוגרים + {} ילדים — נשמר הסה"כ מהקובץ'.format(
                    rno, display, total, adults, kids
                )
            )
        households.append({
            "display_name": display,
            "group_name": group,
            "phone": phone,
            "invited_total": invited,
        })
    return households, report


def import_households(households):
    conn = db()
    added = updated = 0
    for h in households:
        row = None
        if h["phone"]:
            row = conn.execute(
                "SELECT id FROM households WHERE phone = ? AND phone != ''",
                (h["phone"],),
            ).fetchone()
        if row is None:
            row = conn.execute(
                "SELECT id FROM households WHERE display_name = ? AND group_name = ?",
                (h["display_name"], h["group_name"]),
            ).fetchone()
        if row:
            conn.execute(
                """UPDATE households SET phone = COALESCE(NULLIF(?, ''), phone),
                       invited_total = ?, group_name = ?
                   WHERE id = ?""",
                (h["phone"], h["invited_total"], h["group_name"], row["id"]),
            )
            updated += 1
        else:
            conn.execute(
                """INSERT INTO households
                   (token, display_name, group_name, phone, invited_total)
                   VALUES (?, ?, ?, ?, ?)""",
                (new_token(), h["display_name"], h["group_name"],
                 h["phone"], h["invited_total"]),
            )
            added += 1
    conn.commit()
    conn.close()
    return added, updated


# --------------------------------------------------------------- routes ----

@app.route("/")
def guest_page():
    return send_from_directory("static", "guest.html")


@app.route("/admin")
def admin_page():
    return send_from_directory("static", "admin.html")


@app.route("/api/guest")
def api_guest():
    token = request.args.get("g", "")
    conn = db()
    h = conn.execute(
        "SELECT * FROM households WHERE token = ?", (token,)
    ).fetchone()
    if not h:
        conn.close()
        return jsonify({"error": "קישור לא תקין"}), 404
    r = conn.execute(
        "SELECT * FROM responses WHERE household_id = ?", (h["id"],)
    ).fetchone()
    conn.close()
    return jsonify({
        "household": {
            "display_name": h["display_name"],
            "group_name": h["group_name"],
            "invited_total": h["invited_total"],
        },
        "response": dict(r) if r else None,
        "settings": public_settings(),
    })


@app.route("/api/guest/response", methods=["POST"])
def api_guest_response():
    data = request.get_json(force=True)
    token = data.get("token", "")
    attending = data.get("attending", "")
    overnight = data.get("overnight", "")
    try:
        party_count = int(data.get("party_count", 0))
    except (TypeError, ValueError):
        party_count = 0
    if attending not in ("yes", "no", "maybe") or overnight not in ("yes", "no", "maybe", ""):
        return jsonify({"error": "ערך לא תקין"}), 400
    conn = db()
    h = conn.execute(
        "SELECT id FROM households WHERE token = ?", (token,)
    ).fetchone()
    if not h:
        conn.close()
        return jsonify({"error": "קישור לא תקין"}), 404
    conn.execute(
        """INSERT INTO responses (household_id, attending, party_count, overnight, updated_at)
           VALUES (?, ?, ?, ?, ?)
           ON CONFLICT(household_id) DO UPDATE SET
             attending = excluded.attending,
             party_count = excluded.party_count,
             overnight = excluded.overnight,
             updated_at = excluded.updated_at""",
        (h["id"], attending, party_count, overnight,
         datetime.now().isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()
    if overnight in ("yes", "maybe"):
        try:
            maybe_send_packing(h["id"], request.host_url.rstrip("/"))
        except Exception:  # noqa: BLE001 - never fail the guest's save
            pass
    return jsonify({"ok": True})


@app.route("/api/admin/overview")
def api_admin_overview():
    conn = db()
    hs = conn.execute(
        """SELECT h.*, r.attending, r.party_count, r.overnight, r.updated_at
           FROM households h LEFT JOIN responses r ON r.household_id = h.id
           ORDER BY h.group_name, h.display_name"""
    ).fetchall()
    settings = get_settings()
    conn.close()
    households = []
    summary = {
        "households_total": 0, "responded": 0, "pending": 0,
        "yes_households": 0, "yes_people": 0,
        "maybe_households": 0, "maybe_people": 0,
        "no_households": 0, "overnight_people": 0,
        "invited_total": 0,
    }
    for h in hs:
        item = {
            "id": h["id"], "token": h["token"],
            "display_name": h["display_name"], "group_name": h["group_name"],
            "phone": h["phone"], "invited_total": h["invited_total"],
            "attending": h["attending"] or "",
            "party_count": h["party_count"] or 0,
            "overnight": h["overnight"] or "",
            "updated_at": h["updated_at"] or "",
        }
        households.append(item)
        summary["households_total"] += 1
        summary["invited_total"] += h["invited_total"] or 0
        if h["attending"]:
            summary["responded"] += 1
            if h["attending"] == "yes":
                summary["yes_households"] += 1
                summary["yes_people"] += h["party_count"] or 0
            elif h["attending"] == "maybe":
                summary["maybe_households"] += 1
                summary["maybe_people"] += h["party_count"] or 0
            elif h["attending"] == "no":
                summary["no_households"] += 1
            if h["overnight"] in ("yes", "maybe"):
                summary["overnight_people"] += h["party_count"] or 0
        else:
            summary["pending"] += 1
    return jsonify({"summary": summary, "households": households,
                    "settings": settings})


@app.route("/api/admin/households", methods=["POST"])
def api_admin_add():
    data = request.get_json(force=True)
    name = (data.get("display_name") or "").strip()
    if not name:
        return jsonify({"error": "חסר שם"}), 400
    conn = db()
    conn.execute(
        """INSERT INTO households (token, display_name, group_name, phone, invited_total)
           VALUES (?, ?, ?, ?, ?)""",
        (new_token(), name, data.get("group_name", ""),
         normalize_phone(data.get("phone", "")),
         int(data.get("invited_total") or 0)),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/admin/households/<int:hid>", methods=["PUT", "DELETE"])
def api_admin_edit(hid):
    conn = db()
    if request.method == "DELETE":
        conn.execute("DELETE FROM responses WHERE household_id = ?", (hid,))
        conn.execute("DELETE FROM households WHERE id = ?", (hid,))
        conn.commit()
        conn.close()
        return jsonify({"ok": True})
    data = request.get_json(force=True)
    conn.execute(
        """UPDATE households SET display_name = ?, group_name = ?,
               phone = ?, invited_total = ? WHERE id = ?""",
        (data.get("display_name", ""), data.get("group_name", ""),
         normalize_phone(data.get("phone", "")),
         int(data.get("invited_total") or 0), hid),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/admin/settings", methods=["POST"])
def api_admin_settings():
    data = request.get_json(force=True)
    conn = db()
    for key, value in data.items():
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) "
            "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
            (key, str(value)),
        )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/admin/import", methods=["POST"])
def api_admin_import():
    if "file" not in request.files:
        return jsonify({"error": "לא צורף קובץ"}), 400
    f = request.files["file"]
    filename = (f.filename or "").lower()
    rows = []
    try:
        if filename.endswith(".csv"):
            text = f.read().decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text)))
        elif filename.endswith((".xlsx", ".xls")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(f.read()), data_only=True)
            ws = wb.worksheets[0]
            rows = [[c.value for c in row] for row in ws.iter_rows()]
        else:
            return jsonify({"error": "פורמט לא נתמך — יש להעלות xlsx או csv"}), 400
    except Exception as exc:  # noqa: BLE001 - report parse errors to the user
        return jsonify({"error": "קריאת הקובץ נכשלה: {}".format(exc)}), 400
    households, report = parse_rows(rows)
    added, updated = import_households(households)
    report["added"] = added
    report["updated"] = updated
    report["parsed"] = len(households)
    return jsonify(report)


ZAPCLAW_API = "https://api.zapclaw.app/api/v1"

ZAPCLAW_TEMPLATES = [
    {
        "name": "wedding_invitation",
        "category": "UTILITY",
        "language": "he",
        "bodyText": (
            "היי {{1}}, אנחנו מתחתנים! 💍 נשמח שתאשרו הגעה ותראו את כל "
            "הפרטים בקישור האישי שלכם: {{2}} — מחכים לראותכם!"
        ),
        "variableExamples": {
            "1": "משפחת כהן",
            "2": "https://example.com/?g=abc123",
        },
    },
    {
        "name": "wedding_packing_list",
        "category": "UTILITY",
        "language": "he",
        "bodyText": (
            "היי {{1}}, איזה כיף שאתם נשארים לישון איתנו! 🎉 רשימת הציוד "
            "המומלצת ללינה מחכה לכם בקישור האישי שלכם: {{2}} — נתראה שם!"
        ),
        "variableExamples": {
            "1": "משפחת כהן",
            "2": "https://example.com/?g=abc123",
        },
    },
]


def zapclaw_request(settings, path, payload):
    """POST to the ZapClaw API. Returns (ok, detail)."""
    import urllib.request
    import uuid
    body = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        ZAPCLAW_API + path, data=body,
        headers={
            "Content-Type": "application/json",
            "X-API-Key": settings.get("zapclaw_api_key", ""),
            "Idempotency-Key": str(uuid.uuid4()),
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            return True, resp.read().decode("utf-8", "ignore")[:500]
    except Exception as exc:  # noqa: BLE001
        detail = str(exc)
        try:
            detail = exc.read().decode("utf-8", "ignore")[:500] or detail
        except Exception:  # noqa: BLE001
            pass
        return False, detail


def zapclaw_send_template(settings, to_phone, template_name, params):
    payload = {
        "from": settings.get("zapclaw_from", ""),
        "to": to_phone,
        "templateName": template_name,
        "language": "he",
        "components": [{
            "type": "body",
            "parameters": [{"type": "text", "text": p} for p in params],
        }],
    }
    return zapclaw_request(settings, "/send/template", payload)


# -------------------------------------------------------------- relay ----

def _relay_authorized():
    """True when the request carries the shared relay key. An empty
    relay_key setting disables the relay API (every request → 401)."""
    expected = get_settings().get("relay_key", "")
    provided = request.headers.get("X-Relay-Key", "")
    if not expected or not provided:
        return False
    return secrets.compare_digest(provided, expected)


def enqueue_send_job(household_id, kind, to_phone, message, once=False):
    """Queue a send job for the relay worker.

    Dedupe: never create a second job while one for the same
    household+kind is still pending. With once=True (packing list) an
    already-sent job also blocks a new one — the once-only semantics of
    the direct packing path, carried into the jobs table.
    """
    conn = db()
    if once:
        existing = conn.execute(
            "SELECT id FROM send_jobs WHERE household_id = ? AND kind = ? "
            "AND status IN ('pending', 'sent')",
            (household_id, kind),
        ).fetchone()
    else:
        existing = conn.execute(
            "SELECT id FROM send_jobs WHERE household_id = ? AND kind = ? "
            "AND status = 'pending'",
            (household_id, kind),
        ).fetchone()
    if not existing:
        conn.execute(
            "INSERT INTO send_jobs (household_id, kind, to_phone, message, status) "
            "VALUES (?, ?, ?, ?, 'pending')",
            (household_id, kind, to_phone, message),
        )
        conn.commit()
    conn.close()


def render_packing_message(settings, display_name, link):
    """Render the packing-list message text exactly as the direct
    (ZapClaw template) mode produces it, with name + link filled in."""
    name = settings.get("zapclaw_packing_template", "wedding_packing_list")
    body = ""
    for tpl in ZAPCLAW_TEMPLATES:
        if tpl["name"] == name:
            body = tpl["bodyText"]
            break
    if not body:
        body = ZAPCLAW_TEMPLATES[1]["bodyText"]
    return body.replace("{{1}}", display_name).replace("{{2}}", link)


def maybe_send_packing(household_id, base_url):
    """Auto-send the packing list (once per household) when a guest
    answers yes/maybe on the overnight question. ZapClaw mode sends the
    template directly; relay mode queues a fully rendered send job for
    the external worker instead."""
    settings = get_settings()
    if settings.get("whatsapp_provider") == "relay":
        conn = db()
        h = conn.execute(
            "SELECT * FROM households WHERE id = ?", (household_id,)
        ).fetchone()
        already = conn.execute(
            "SELECT id FROM send_log WHERE household_id = ? "
            "AND channel = 'relay_packing' AND status = 'sent'",
            (household_id,),
        ).fetchone()
        conn.close()
        if not h or not h["phone"] or already:
            return
        link = "{}?g={}".format(base_url, h["token"])
        message = render_packing_message(
            settings, h["display_name"], link)
        enqueue_send_job(household_id, "packing", h["phone"], message,
                         once=True)
        return
    if settings.get("whatsapp_provider") != "zapclaw":
        return
    if not settings.get("zapclaw_api_key") or not settings.get("zapclaw_from"):
        return
    conn = db()
    h = conn.execute(
        "SELECT * FROM households WHERE id = ?", (household_id,)
    ).fetchone()
    already = conn.execute(
        "SELECT id FROM send_log WHERE household_id = ? "
        "AND channel = 'zapclaw_packing' AND status = 'sent'",
        (household_id,),
    ).fetchone()
    conn.close()
    if not h or not h["phone"] or already:
        return
    link = "{}?g={}".format(base_url, h["token"])
    ok, detail = zapclaw_send_template(
        settings, h["phone"],
        settings.get("zapclaw_packing_template", "wedding_packing_list"),
        [h["display_name"], link],
    )
    conn = db()
    conn.execute(
        "INSERT INTO send_log (household_id, channel, status, detail) "
        "VALUES (?, ?, ?, ?)",
        (household_id, "zapclaw_packing", "sent" if ok else "failed", detail),
    )
    conn.commit()
    conn.close()


@app.route("/api/admin/zapclaw/create-templates", methods=["POST"])
def api_admin_zapclaw_templates():
    """Create the two Hebrew message templates in ZapClaw/Meta and submit
    them for approval. Requires the API key + sender number in settings."""
    settings = get_settings()
    if not settings.get("zapclaw_api_key") or not settings.get("zapclaw_from"):
        return jsonify({
            "error": "יש למלא קודם מפתח API ומספר שולח של ZapClaw ולשמור הגדרות"
        }), 400
    results = []
    for tpl in ZAPCLAW_TEMPLATES:
        payload = dict(tpl)
        payload["from"] = settings["zapclaw_from"]
        ok, detail = zapclaw_request(settings, "/send/templates", payload)
        results.append({"name": tpl["name"], "ok": ok, "detail": detail})
    return jsonify({"ok": all(r["ok"] for r in results), "results": results})


@app.route("/api/admin/send", methods=["POST"])
def api_admin_send():
    """Send (or prepare) a WhatsApp invitation for one household.

    Modes (settings.whatsapp_provider):
    - wa_link (default): returns a wa.me link; the admin taps it and sends
      from their own WhatsApp in one tap.
    - gateway: POSTs {phone, message} to a configured HTTP gateway
      (WAHA / Evolution API / Ultramsg-style) and sends automatically.
    - relay: queues a send job in send_jobs; an external worker pulls
      it via /api/relay/jobs, sends it, and posts the result back.
    """
    data = request.get_json(force=True)
    hid = data.get("household_id")
    settings = get_settings()
    conn = db()
    h = conn.execute("SELECT * FROM households WHERE id = ?", (hid,)).fetchone()
    conn.close()
    if not h:
        return jsonify({"error": "בית אב לא נמצא"}), 404
    if not h["phone"]:
        return jsonify({"error": "לבית אב זה אין מספר טלפון"}), 400
    base = request.host_url.rstrip("/")
    link = "{}?g={}".format(base, h["token"])
    message = settings.get("invite_template", "").replace(
        "{name}", h["display_name"]).replace("{link}", link)
    provider = settings.get("whatsapp_provider", "wa_link")
    if provider == "relay":
        enqueue_send_job(hid, "invite", h["phone"], message)
        return jsonify({"ok": True, "queued": True, "mode": "relay"})
    if provider == "zapclaw" and settings.get("zapclaw_api_key") \
            and settings.get("zapclaw_from"):
        ok, detail = zapclaw_send_template(
            settings, h["phone"],
            settings.get("zapclaw_invite_template", "wedding_invitation"),
            [h["display_name"], link],
        )
        status = "sent" if ok else "failed"
        conn = db()
        conn.execute(
            "INSERT INTO send_log (household_id, channel, status, detail) VALUES (?, ?, ?, ?)",
            (hid, "zapclaw", status, detail),
        )
        conn.commit()
        conn.close()
        return jsonify({"ok": ok, "mode": "zapclaw",
                        "status": status, "detail": detail})
    if provider == "gateway" and settings.get("gateway_url"):
        import urllib.request
        payload = json.dumps(
            {"phone": h["phone"], "message": message}
        ).encode("utf-8")
        req = urllib.request.Request(
            settings["gateway_url"], data=payload,
            headers={"Content-Type": "application/json",
                     "Authorization": "Bearer {}".format(settings.get("gateway_token", ""))},
            method="POST",
        )
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                detail = resp.read().decode("utf-8", "ignore")[:500]
            status = "sent"
        except Exception as exc:  # noqa: BLE001
            detail, status = str(exc), "failed"
        conn = db()
        conn.execute(
            "INSERT INTO send_log (household_id, channel, status, detail) VALUES (?, ?, ?, ?)",
            (hid, "gateway", status, detail),
        )
        conn.commit()
        conn.close()
        return jsonify({"ok": status == "sent", "mode": "gateway",
                        "status": status, "detail": detail})
    url = wa_link(h["phone"], message)
    conn = db()
    conn.execute(
        "INSERT INTO send_log (household_id, channel, status, detail) VALUES (?, ?, ?, ?)",
        (hid, "wa_link", "prepared", url),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "mode": "wa_link", "url": url})


@app.route("/api/relay/jobs")
def api_relay_jobs():
    """Pending send jobs for the external relay worker, oldest first."""
    if not _relay_authorized():
        return jsonify({"error": "unauthorized"}), 401
    conn = db()
    rows = conn.execute(
        "SELECT id, household_id, kind, to_phone, message FROM send_jobs "
        "WHERE status = 'pending' ORDER BY id ASC LIMIT 20"
    ).fetchall()
    conn.close()
    return jsonify({"jobs": [{
        "id": r["id"], "household_id": r["household_id"],
        "kind": r["kind"], "to": r["to_phone"], "message": r["message"],
    } for r in rows]})


@app.route("/api/relay/jobs/<int:job_id>/result", methods=["POST"])
def api_relay_job_result(job_id):
    """Record a relay worker's send outcome for one job.

    Bookkeeping mirrors the direct send paths in the messages log
    (send_log): invites log under channel 'relay', packing under
    'relay_packing' — whose 'sent' rows are also the packing once-only
    guard, exactly like the 'zapclaw_packing' rows in direct mode. The
    send_log row (with its created_at) is the invite/packing "sent at"
    record, same as in the direct paths, which keep no separate
    timestamp column either.
    """
    if not _relay_authorized():
        return jsonify({"error": "unauthorized"}), 401
    data = request.get_json(silent=True) or {}
    status = data.get("status", "")
    if status not in ("sent", "failed"):
        return jsonify({"error": "סטטוס לא תקין"}), 400
    provider_message_id = data.get("provider_message_id") or None
    error = data.get("error") or None
    conn = db()
    job = conn.execute(
        "SELECT * FROM send_jobs WHERE id = ?", (job_id,)
    ).fetchone()
    if not job:
        conn.close()
        return jsonify({"error": "משימת שליחה לא נמצאה"}), 404
    was_pending = job["status"] == "pending"
    conn.execute(
        "UPDATE send_jobs SET status = ?, provider_message_id = ?, "
        "error = ?, updated_at = ? WHERE id = ?",
        (status, provider_message_id, error,
         datetime.now().isoformat(timespec="seconds"), job_id),
    )
    if was_pending:
        channel = "relay" if job["kind"] == "invite" else "relay_packing"
        detail = provider_message_id if status == "sent" else (error or "")
        conn.execute(
            "INSERT INTO send_log (household_id, channel, status, detail) "
            "VALUES (?, ?, ?, ?)",
            (job["household_id"], channel, status, detail),
        )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/admin/export")
def api_admin_export():
    conn = db()
    rows = conn.execute(
        """SELECT h.display_name, h.group_name, h.phone, h.invited_total,
                  r.attending, r.party_count, r.overnight, r.updated_at
           FROM households h LEFT JOIN responses r ON r.household_id = h.id
           ORDER BY h.group_name, h.display_name"""
    ).fetchall()
    conn.close()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["שם בית אב", "קבוצה", "טלפון", "מוזמנים", "מגיע",
                "כמה מגיעים", "לינה", "עודכן"])
    for r in rows:
        w.writerow([r[0], r[1], r[2], r[3], r[4] or "", r[5] or "",
                    r[6] or "", r[7] or ""])
    data = buf.getvalue().encode("utf-8-sig")
    return app.response_class(
        data, mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=rsvp-export.csv"},
    )


if __name__ == "__main__":
    init_db()
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port)
else:
    init_db()
